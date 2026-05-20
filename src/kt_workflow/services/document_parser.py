"""从本地文件抽取纯文本（PDF / Word / PPT / Excel / 纯文本等）。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import chardet


def _truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars] + "\n\n…（已截断）", True


def _read_plain_bytes(data: bytes) -> str:
    detected = chardet.detect(data)
    encoding = detected.get("encoding") if detected else None
    if encoding:
        return data.decode(encoding, errors="replace")
    return data.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> str:
    import pypdf

    reader = pypdf.PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text)
    return "\n\n".join(parts)


def _parse_docx(data: bytes) -> str:
    from docx2python import docx2python

    doc_result = docx2python(BytesIO(data))
    all_parts: list[str] = []
    for section in doc_result.body:
        if not isinstance(section, list):
            continue
        for item in section:
            if isinstance(item, list):
                for sub_item in item:
                    if isinstance(sub_item, str) and sub_item.strip():
                        all_parts.append(sub_item.strip())
                    elif isinstance(sub_item, list):
                        row_text = "\n".join(
                            str(cell).strip() for cell in sub_item if str(cell).strip()
                        )
                        if row_text:
                            all_parts.append(row_text)
            elif isinstance(item, str) and item.strip():
                all_parts.append(item.strip())
    doc_result.close()
    return "\n\n".join(all_parts)


def _parse_excel(data: bytes, ext: str) -> str:
    import pandas as pd

    stream = BytesIO(data)
    if ext == ".csv":
        df = pd.read_csv(stream)
    else:
        df = pd.read_excel(stream)
    return df.to_string()


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    full_text: list[str] = []
    for i, slide in enumerate(prs.slides):
        page_content = [f"=== 第 {i + 1} 页 ==="]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                page_content.append(shape.text.strip())
            if shape.has_table:
                table_texts = []
                for row in shape.table.rows:
                    row_cells = [
                        cell.text_frame.text.strip()
                        for cell in row.cells
                        if cell.text_frame.text.strip()
                    ]
                    if row_cells:
                        table_texts.append(" | ".join(row_cells))
                if table_texts:
                    page_content.append("[表格]\n" + "\n".join(table_texts))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                page_content.append(f"[备注]: {notes.strip()}")
        full_text.append("\n".join(page_content))
    return "\n\n".join(full_text)


def extract_text_from_path(path: Path, max_chars: int = 50_000) -> tuple[str, dict[str, Any]]:
    """读取本地文件并返回 (正文, meta)。"""
    ext = path.suffix.lower()
    meta: dict[str, Any] = {
        "source_path": str(path),
        "file_format": ext.lstrip(".") or "unknown",
        "extractor": "document_parser_v1",
    }

    if not path.is_file():
        msg = f"[extract] 路径不存在或非文件: {path}"
        meta["error"] = "not_found"
        return msg, meta

    try:
        data = path.read_bytes()
    except OSError as exc:
        msg = f"[extract] 无法读取文件: {path} ({exc})"
        meta["error"] = "read_failed"
        return msg, meta

    meta["byte_size"] = len(data)

    try:
        if ext == ".pdf":
            text = _parse_pdf(data)
        elif ext in (".docx", ".doc"):
            text = _parse_docx(data)
        elif ext in (".xlsx", ".xls", ".csv"):
            text = _parse_excel(data, ext)
        elif ext in (".pptx", ".ppt"):
            text = _parse_pptx(data)
        elif ext in (".txt", ".md", ".json", ".xml", ".html", ".htm"):
            text = _read_plain_bytes(data)
        else:
            text = _read_plain_bytes(data)
            meta["note"] = f"未识别扩展名 {ext}，按纯文本尝试解码"
    except Exception as exc:
        text = f"[extract] 解析失败 ({ext}): {exc}"
        meta["error"] = "parse_failed"
        meta["parse_error"] = str(exc)

    text = text.strip()
    meta["char_count_raw"] = len(text)
    text, truncated = _truncate(text, max_chars)
    meta["char_count"] = len(text)
    meta["max_chars"] = max_chars
    meta["truncated"] = truncated
    return text, meta
